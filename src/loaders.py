"""Leitura e validação estrutural de arquivos de entrada.

Formatos aceitos:
- cadastro: XLSX ou CSV;
- atividade: XLSX/CSV exportados do Moodle ou o texto copiado da página
  "Participantes" do Moodle (TXT);
- base escolar: ZIP/CSV dos microdados do Censo Escolar ou XLSX/CSV
  exportado do Catálogo de Escolas (ver ``school_data``).

As mensagens de erro nunca reproduzem conteúdo de células: citam apenas
arquivo, linha e nome de coluna.
"""

from __future__ import annotations

import csv
import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime

import pandas as pd

from .normalization import chave_cabecalho, eh_vazio, limpar_texto

LIMITE_PLANILHA_BYTES = 25 * 1024 * 1024
LIMITE_BASE_ESCOLAR_BYTES = 120 * 1024 * 1024
LIMITE_DESCOMPACTADO_XLSX = 300 * 1024 * 1024
LIMITE_LINHAS = 200_000
LIMITE_COLUNAS = 400
MAX_LINHAS_BUSCA_CABECALHO = 30


class ErroImportacao(Exception):
    """Erro com mensagem segura para exibir ao usuário (sem dados pessoais)."""


# ---------------------------------------------------------------- dicionário de colunas


@dataclass(frozen=True)
class Campo:
    nome: str
    rotulo: str
    aliases: tuple[str, ...]
    obrigatorio: bool = False
    recomendado: bool = False


def definir_campo(nome: str, rotulo: str, *aliases: str, obrigatorio=False, recomendado=False) -> Campo:
    chaves = tuple(dict.fromkeys(chave_cabecalho(a) for a in (rotulo, *aliases)))
    return Campo(nome, rotulo, chaves, obrigatorio, recomendado)


# Cabeçalhos originais -> nomes internos. Comparação sem acento, caixa e pontuação.
CAMPOS_CADASTRO: tuple[Campo, ...] = (
    definir_campo("regiao", "Região", "regiao geografica", recomendado=True),
    definir_campo("uf", "UF", "sigla uf", "sigla da uf", "unidade da federacao", recomendado=True),
    definir_campo("ente", "Ente", "ente federado", "rede de ensino", "secretaria", recomendado=True),
    definir_campo("cod_ibge", "CODIBGE", "cod ibge", "codigo ibge", "codigo do ibge", "ibge",
           "codigo ibge do ente", "codigo do municipio", recomendado=True),
    definir_campo("data_resposta", "Data resposta", "data da resposta", "data de resposta",
           "data inscricao", "data da inscricao"),
    definir_campo("ciclo", "Ciclo", recomendado=True),
    definir_campo("articulador", "Articulador", "articuladora", "articulador a"),
    definir_campo("esfera", "Esfera", "esfera administrativa"),
    definir_campo("nome", "Nome", "nome completo", "nome do cursista", "cursista", obrigatorio=True),
    definir_campo("cpf", "CPF", "numero do cpf"),
    definir_campo("data_nascimento", "Data Nascimento", "data de nascimento", "nascimento", "dt nascimento"),
    definir_campo("email", "E-mail", "email", "endereco de email", "endereco de e mail",
           "correio eletronico", obrigatorio=True),
    definir_campo("telefone", "Telefone", "celular", "fone", "telefone celular"),
    definir_campo("whatsapp", "WhatsApp", "whats app", "possui whatsapp"),
    definir_campo("tipo", "Tipo"),
    definir_campo("origem", "Origem"),
    definir_campo("situacao", "Situação", "status", "situacao da inscricao", recomendado=True),
    definir_campo("raca_cor", "Raça/Cor", "raca cor", "cor raca"),
    definir_campo("genero", "Gênero", "sexo"),
    definir_campo("escolaridade", "Escolaridade", "formacao", "nivel de escolaridade"),
    definir_campo("cargo_funcao", "Cargo/Função", "cargo", "funcao", "cargo funcao"),
    definir_campo("escola", "Escola", "nome da escola", "unidade escolar", recomendado=True),
    definir_campo("inep", "INEP", "codigo inep", "cod inep", "inep da escola", "codigo da escola",
           "co entidade", recomendado=True),
    definir_campo("id_compartilhado", "Identificador compartilhado", "id compartilhado"),
)

CAMPOS_ATIVIDADE: tuple[Campo, ...] = (
    definir_campo("nome_completo", "Nome completo", "nome sobrenome", "full name", "participante"),
    definir_campo("nome", "Nome", "primeiro nome", "first name"),
    definir_campo("sobrenome", "Sobrenome", "ultimo nome", "last name", "surname"),
    definir_campo("email", "Endereço de email", "email", "e mail", "endereco de e mail",
           "email address", obrigatorio=True),
    definir_campo("papeis", "Papéis", "papel", "roles", "role", recomendado=True),
    definir_campo("grupos", "Grupos", "grupo", "groups", "group", recomendado=True),
    definir_campo("ultimo_acesso", "Último acesso ao curso", "ultimo acesso", "last access to course",
           "last access", obrigatorio=True),
    definir_campo("situacao_conta", "Situação", "status", "situacao da inscricao", recomendado=True),
    definir_campo("id_compartilhado", "Identificador compartilhado", "id compartilhado",
           "numero de identificacao", "id number", "idnumber"),
)


@dataclass
class Mapeamento:
    original_para_interno: dict[str, str]
    faltando_obrigatorios: list[str] = field(default_factory=list)
    faltando_recomendados: list[str] = field(default_factory=list)
    nao_mapeadas: list[str] = field(default_factory=list)
    ambiguidades: list[str] = field(default_factory=list)

    @property
    def interno_para_original(self) -> dict[str, str]:
        return {v: k for k, v in self.original_para_interno.items()}


def mapear_colunas(
    cabecalhos: list[str], campos: tuple[Campo, ...], manual: dict[str, str] | None = None
) -> Mapeamento:
    """Associa cada cabeçalho original a um campo interno.

    ``manual`` (campo interno -> cabeçalho original) tem prioridade e permite
    corrigir na interface o que a correspondência automática não encontrou.
    """
    manual = {k: v for k, v in (manual or {}).items() if v in cabecalhos}
    usados = set(manual.values())
    resultado: dict[str, str] = {orig: interno for interno, orig in manual.items()}
    ambiguidades = []
    chaves = {h: chave_cabecalho(h) for h in cabecalhos}
    for campo in campos:
        if campo.nome in manual:
            continue
        candidatos = [h for h in cabecalhos if h not in usados and chaves[h] in campo.aliases]
        if not candidatos:
            continue
        if len(candidatos) > 1:
            ambiguidades.append(
                f"Mais de uma coluna pode ser “{campo.rotulo}”: "
                + ", ".join(f"“{c}”" for c in candidatos)
                + f". Foi usada “{candidatos[0]}”."
            )
        resultado[candidatos[0]] = campo.nome
        usados.add(candidatos[0])
    internos = set(resultado.values())
    return Mapeamento(
        original_para_interno=resultado,
        faltando_obrigatorios=[c.rotulo for c in campos if c.obrigatorio and c.nome not in internos],
        faltando_recomendados=[c.rotulo for c in campos if c.recomendado and c.nome not in internos],
        nao_mapeadas=[h for h in cabecalhos if h not in resultado],
        ambiguidades=ambiguidades,
    )


# ---------------------------------------------------------------- arquivos


@dataclass
class TabelaBruta:
    """Conteúdo de um arquivo, com os valores das células preservados como vieram."""

    arquivo: str
    aba: str | None
    formato: str
    sha256: str
    tamanho: int
    linha_cabecalho: int
    cabecalhos: list[str]
    dados: pd.DataFrame  # dtype object; coluna "__linha__" = linha no arquivo original
    avisos: list[str] = field(default_factory=list)

    def metadados(self) -> dict:
        return {
            "arquivo": self.arquivo,
            "aba": self.aba,
            "formato": self.formato,
            "sha256": self.sha256,
            "tamanho_bytes": self.tamanho,
            "linha_cabecalho": self.linha_cabecalho,
            "linhas": int(len(self.dados)),
            "cabecalhos": list(self.cabecalhos),
        }


def identificar_formato(nome: str, conteudo: bytes, permitidos: tuple[str, ...]) -> str:
    """Confere extensão, tamanho e assinatura binária. Retorna xlsx|csv|txt|zip."""
    extensao = nome.rsplit(".", 1)[-1].lower() if "." in nome else ""
    if extensao not in permitidos:
        raise ErroImportacao(
            f"O arquivo “{nome}” tem extensão .{extensao or '?'}; são aceitos: "
            + ", ".join(f".{p}" for p in permitidos) + "."
        )
    limite = LIMITE_BASE_ESCOLAR_BYTES if extensao == "zip" else LIMITE_PLANILHA_BYTES
    if len(conteudo) == 0:
        raise ErroImportacao(f"O arquivo “{nome}” está vazio.")
    if len(conteudo) > limite:
        raise ErroImportacao(
            f"O arquivo “{nome}” tem {len(conteudo) / 1048576:.1f} MB; o limite é {limite // 1048576} MB."
        )
    if extensao in ("xlsx", "zip"):
        if not conteudo.startswith(b"PK\x03\x04"):
            raise ErroImportacao(f"“{nome}” não é um arquivo .{extensao} válido (assinatura incorreta).")
        try:
            with zipfile.ZipFile(io.BytesIO(conteudo)) as zf:
                infos = zf.infolist()
        except zipfile.BadZipFile as exc:
            raise ErroImportacao(f"“{nome}” está corrompido ou não é um arquivo compactado válido.") from exc
        if extensao == "xlsx":
            nomes = {i.filename for i in infos}
            if "xl/workbook.xml" not in nomes:
                raise ErroImportacao(f"“{nome}” não contém uma pasta de trabalho do Excel.")
            if any(n.lower().endswith("vbaproject.bin") for n in nomes):
                raise ErroImportacao(f"“{nome}” contém macros; salve como .xlsx sem macros.")
            if sum(i.file_size for i in infos) > LIMITE_DESCOMPACTADO_XLSX:
                raise ErroImportacao(f"“{nome}” é grande demais depois de descompactado.")
    else:
        if b"\x00" in conteudo[:65536]:
            raise ErroImportacao(f"“{nome}” não parece ser um arquivo de texto.")
    return extensao


def decodificar_texto(conteudo: bytes) -> str:
    for codificacao in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return conteudo.decode(codificacao)
        except UnicodeDecodeError:
            continue
    raise ErroImportacao("Não foi possível identificar a codificação do texto.")  # pragma: no cover


def listar_abas(conteudo: bytes) -> list[str]:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as exc:
        raise ErroImportacao("Não foi possível abrir a planilha (arquivo danificado ou protegido).") from exc
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _linhas_xlsx(conteudo: bytes, aba: str | None) -> tuple[list[list], str]:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception as exc:
        raise ErroImportacao("Não foi possível abrir a planilha (arquivo danificado ou protegido).") from exc
    try:
        if aba is not None and aba not in wb.sheetnames:
            raise ErroImportacao("A aba selecionada não existe na planilha.")
        ws = wb[aba] if aba else wb.worksheets[0]
        linhas = []
        for i, linha in enumerate(ws.iter_rows(values_only=True)):
            if i > LIMITE_LINHAS + MAX_LINHAS_BUSCA_CABECALHO:
                raise ErroImportacao(f"A planilha excede o limite de {LIMITE_LINHAS:,} linhas.".replace(",", "."))
            linhas.append(list(linha[:LIMITE_COLUNAS]))
        return linhas, ws.title
    finally:
        wb.close()


def _linhas_csv(texto: str) -> list[list]:
    amostra = texto[:20000]
    try:
        dialeto = csv.Sniffer().sniff(amostra, delimiters=";,\t|")
        delimitador = dialeto.delimiter
    except csv.Error:
        contagens = {d: amostra.count(d) for d in ";,\t|"}
        delimitador = max(contagens, key=contagens.get)
    leitor = csv.reader(io.StringIO(texto), delimiter=delimitador)
    linhas = []
    for i, linha in enumerate(leitor):
        if i > LIMITE_LINHAS + MAX_LINHAS_BUSCA_CABECALHO:
            raise ErroImportacao(f"O arquivo excede o limite de {LIMITE_LINHAS:,} linhas.".replace(",", "."))
        linhas.append(linha[:LIMITE_COLUNAS])
    return linhas


RE_EMAIL_SOLTO = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
CABECALHOS_MOODLE_TEXTO = [
    "Nome completo", "Endereço de email", "Papéis", "Grupos", "Último acesso ao curso", "Situação",
]


def ler_texto_participantes_moodle(texto: str) -> tuple[list[list], list[int], list[int]]:
    """Lê o texto copiado da página "Participantes" do Moodle.

    Cada participante ocupa três linhas: ``Selecione 'NOME'``; a linha de dados
    separada por TAB (nome, e-mail, papéis, grupos, último acesso); e a
    situação da inscrição. Retorna (registros, linhas_de_origem,
    linhas_não_reconhecidas), em que as não reconhecidas são linhas com
    estrutura de dados (4 ou mais TABs) mas sem e-mail identificável.
    """
    linhas = texto.splitlines()
    registros, origem, nao_reconhecidas = [], [], []
    for i, linha in enumerate(linhas):
        campos = [c.strip() for c in linha.split("\t")]
        posicoes = [j for j, c in enumerate(campos) if RE_EMAIL_SOLTO.match(c)]
        if not posicoes:
            if linha.count("\t") >= 4 and sum(1 for c in campos if c) >= 3:
                nao_reconhecidas.append(i + 1)
            continue
        j = posicoes[0]
        nome = next((c for c in reversed(campos[:j]) if c), None)
        resto = campos[j + 1:] + [None, None, None]
        situacao = None
        for k in range(i + 1, min(i + 3, len(linhas))):
            seguinte = linhas[k].strip()
            if not seguinte:
                continue
            if ("\t" not in linhas[k] and "@" not in seguinte and len(seguinte.split()) <= 3
                    and not seguinte.lower().startswith(("selecione", "select"))):
                situacao = seguinte
            break
        registros.append([nome, campos[j], resto[0] or None, resto[1] or None, resto[2] or None, situacao])
        origem.append(i + 1)
    return registros, origem, nao_reconhecidas


def detectar_linha_cabecalho(linhas: list[list], campos: tuple[Campo, ...]) -> int:
    """Índice (base 0) da linha com mais nomes de coluna reconhecidos."""
    aliases = {a: c.nome for c in campos for a in c.aliases}
    melhor, melhor_pontos = 0, 1
    for i, linha in enumerate(linhas[:MAX_LINHAS_BUSCA_CABECALHO]):
        reconhecidos = {aliases[k] for k in (chave_cabecalho(v) for v in linha) if k in aliases}
        if len(reconhecidos) > melhor_pontos:
            melhor, melhor_pontos = i, len(reconhecidos)
    return melhor


def _nomes_unicos(brutos: list) -> list[str]:
    nomes, vistos = [], {}
    for j, valor in enumerate(brutos):
        nome = limpar_texto(valor) or f"(coluna {j + 1} sem título)"
        if nome in vistos:
            vistos[nome] += 1
            nome = f"{nome} ({vistos[nome]})"
        else:
            vistos[nome] = 1
        nomes.append(nome)
    return nomes


def _valor_celula(valor):
    if isinstance(valor, str):
        return valor if valor.strip() else None
    if isinstance(valor, (datetime, date)) or valor is None:
        return valor
    if isinstance(valor, float) and valor != valor:  # NaN
        return None
    return valor


def ler_tabela(
    nome: str,
    conteudo: bytes,
    campos: tuple[Campo, ...],
    permitidos: tuple[str, ...] = ("xlsx", "csv"),
    aba: str | None = None,
) -> TabelaBruta:
    formato = identificar_formato(nome, conteudo, permitidos)
    avisos: list[str] = []
    if formato == "xlsx":
        linhas, aba = _linhas_xlsx(conteudo, aba)
        indice_cab = detectar_linha_cabecalho(linhas, campos)
        cabecalhos = _nomes_unicos(linhas[indice_cab] if linhas else [])
        corpo = linhas[indice_cab + 1:]
        primeira_linha_dados = indice_cab + 2
    else:
        texto = decodificar_texto(conteudo)
        registros, origem, nao_rec = (ler_texto_participantes_moodle(texto) if formato == "txt" else ([], [], []))
        if registros:
            cabecalhos = list(CABECALHOS_MOODLE_TEXTO)
            dados = pd.DataFrame(registros, columns=cabecalhos, dtype=object)
            dados["__linha__"] = origem
            if nao_rec:
                avisos.append(
                    f"{len(nao_rec)} linha(s) com estrutura de participante sem e-mail reconhecível "
                    f"não foram lidas (linhas {', '.join(map(str, nao_rec[:10]))}"
                    + ("…" if len(nao_rec) > 10 else "") + ")."
                )
            return TabelaBruta(nome, None, "texto do Moodle", hashlib.sha256(conteudo).hexdigest(),
                               len(conteudo), 0, cabecalhos, dados, avisos)
        linhas = _linhas_csv(texto)
        indice_cab = detectar_linha_cabecalho(linhas, campos)
        cabecalhos = _nomes_unicos(linhas[indice_cab] if linhas else [])
        corpo = linhas[indice_cab + 1:]
        primeira_linha_dados = indice_cab + 2
    if not cabecalhos:
        raise ErroImportacao(f"“{nome}” não tem linhas.")
    largura = len(cabecalhos)
    registros, origem, extras = [], [], 0
    for deslocamento, linha in enumerate(corpo):
        valores = [_valor_celula(v) for v in linha]
        if all(eh_vazio(v) for v in valores):
            continue
        if any(not eh_vazio(v) for v in valores[largura:]):
            extras += 1
        valores = (valores + [None] * largura)[:largura]
        registros.append(valores)
        origem.append(primeira_linha_dados + deslocamento)
    if len(registros) > LIMITE_LINHAS:
        raise ErroImportacao(f"“{nome}” excede o limite de {LIMITE_LINHAS:,} registros.".replace(",", "."))
    if extras:
        avisos.append(f"{extras} linha(s) têm conteúdo em colunas sem cabeçalho; esse conteúdo foi ignorado.")
    dados = pd.DataFrame(registros, columns=cabecalhos, dtype=object)
    dados["__linha__"] = origem
    return TabelaBruta(nome, aba, formato, hashlib.sha256(conteudo).hexdigest(), len(conteudo),
                       indice_cab + 1, cabecalhos, dados, avisos)
