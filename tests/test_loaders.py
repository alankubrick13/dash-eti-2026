"""Leitura dos arquivos reais: planilhas do cadastro, texto do Moodle e recusa de formatos incompatíveis.

As variações (CSV, extensão errada, pacote com macro) são feitas a partir dos próprios arquivos reais.
"""

import io
import zipfile

import openpyxl
import pytest
from conftest import confere, valores

from src.loaders import (
    CAMPOS_ATIVIDADE,
    CAMPOS_CADASTRO,
    ErroImportacao,
    identificar_formato,
    ler_tabela,
    ler_texto_participantes_moodle,
    listar_abas,
    mapear_colunas,
)
from src.normalization import eh_vazio, validar_cpf


def _linhas_da_planilha(conteudo: bytes, aba: str) -> list[tuple]:
    planilha = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    return [linha for linha in planilha[aba].iter_rows(values_only=True)]


def test_planilhas_reais_lidas_com_todas_as_colunas(arquivos_cadastro, tabelas_cadastro):
    for arquivo, tabela in zip(arquivos_cadastro, tabelas_cadastro):
        mapa = mapear_colunas(tabela.cabecalhos, CAMPOS_CADASTRO)
        confere(not mapa.faltando_obrigatorios and not mapa.nao_mapeadas, f"{arquivo.name}: coluna não reconhecida")
        linhas = _linhas_da_planilha(arquivo.read_bytes(), tabela.aba)
        com_dados = [l for l in linhas[tabela.linha_cabecalho:] if any(not eh_vazio(v) for v in l)]
        confere(len(tabela.dados) == len(com_dados), f"{arquivo.name}: número de linhas difere da planilha")
        numeros = tabela.dados["__linha__"].astype(int).tolist()
        confere(numeros == sorted(numeros) and numeros[0] > tabela.linha_cabecalho,
                f"{arquivo.name}: número de linha original não preservado")


def test_valores_das_celulas_preservados_como_vieram(arquivos_cadastro, tabelas_cadastro):
    divergencias = 0
    for arquivo, tabela in zip(arquivos_cadastro, tabelas_cadastro):
        linhas = _linhas_da_planilha(arquivo.read_bytes(), tabela.aba)
        for coluna in ("INEP", "CPF", "CODIBGE"):
            if coluna not in tabela.cabecalhos:
                continue
            j = tabela.cabecalhos.index(coluna)
            for _, registro in tabela.dados.iterrows():
                original = linhas[int(registro["__linha__"]) - 1][j]
                lido = registro[coluna]
                if not (eh_vazio(original) and eh_vazio(lido)) and str(original).strip() != str(lido).strip():
                    divergencias += 1
    confere(divergencias == 0, f"{divergencias} célula(s) real(is) alterada(s) na leitura")


def test_abas_das_planilhas_reais(arquivos_cadastro):
    for arquivo in arquivos_cadastro:
        nomes = openpyxl.load_workbook(io.BytesIO(arquivo.read_bytes()), read_only=True).sheetnames
        assert listar_abas(arquivo.read_bytes()) == nomes


def test_cadastro_real_em_csv(arquivos_cadastro, tabelas_cadastro):
    """A mesma planilha exportada em CSV (ponto e vírgula, cp1252) dá o mesmo conteúdo."""
    for arquivo, tabela in zip(arquivos_cadastro, tabelas_cadastro):
        linhas = [";".join("" if eh_vazio(v) else str(v) for v in [*tabela.cabecalhos])]
        for _, registro in tabela.dados.iterrows():
            linhas.append(";".join("" if eh_vazio(registro[c]) else str(registro[c]) for c in tabela.cabecalhos))
        texto = "\n".join(linhas) + "\n"
        try:
            conteudo = texto.encode("cp1252")
        except UnicodeEncodeError:
            conteudo = texto.encode("utf-8")
        csv = ler_tabela(arquivo.with_suffix(".csv").name, conteudo, CAMPOS_CADASTRO)
        assert csv.cabecalhos == tabela.cabecalhos
        assert len(csv.dados) == len(tabela.dados)
        inep_xlsx = ["" if eh_vazio(v) else str(v) for v in tabela.dados["INEP"]]
        inep_csv = ["" if eh_vazio(v) else str(v) for v in csv.dados["INEP"]]
        confere(inep_xlsx == inep_csv, f"{arquivo.name}: códigos INEP diferem entre XLSX e CSV")


def test_texto_real_da_pagina_de_participantes(arquivo_atividade, tabela_atividade):
    texto = arquivo_atividade.read_text(encoding="utf-8", errors="replace")
    registros, _, nao_reconhecidas = ler_texto_participantes_moodle(texto)
    blocos = texto.count("Selecione '")
    confere(len(registros) + len(nao_reconhecidas) == blocos, "participante do texto real descartado em silêncio")
    confere(all(len(r) == 6 and "@" in (r[1] or "") for r in registros), "registro do Moodle lido incompleto")
    assert tabela_atividade.formato == "texto do Moodle"
    assert len(tabela_atividade.dados) == len(registros)
    assert not mapear_colunas(tabela_atividade.cabecalhos, CAMPOS_ATIVIDADE).faltando_obrigatorios


def test_mapeamento_manual_com_os_cabecalhos_reais(tabelas_cadastro):
    cabecalhos = tabelas_cadastro[0].cabecalhos
    coluna_email = next(c for c, i in mapear_colunas(cabecalhos, CAMPOS_CADASTRO).original_para_interno.items()
                        if i == "email")
    renomeados = [f"{c} (coluna)" if c == coluna_email else c for c in cabecalhos]
    assert mapear_colunas(renomeados, CAMPOS_CADASTRO).faltando_obrigatorios == ["E-mail"]
    manual = mapear_colunas(renomeados, CAMPOS_CADASTRO, manual={"email": f"{coluna_email} (coluna)"})
    assert not manual.faltando_obrigatorios


def test_recusa_formatos_incompativeis(arquivos_cadastro, arquivo_atividade):
    planilha = arquivos_cadastro[0].read_bytes()
    with pytest.raises(ErroImportacao, match="extensão"):
        identificar_formato(arquivos_cadastro[0].with_suffix(".xls").name, planilha, ("xlsx", "csv"))
    with pytest.raises(ErroImportacao, match="assinatura"):
        identificar_formato(arquivo_atividade.with_suffix(".xlsx").name, arquivo_atividade.read_bytes(), ("xlsx",))
    with pytest.raises(ErroImportacao, match="vazio"):
        identificar_formato("cadastro.csv", b"", ("csv",))
    with pytest.raises(ErroImportacao, match="texto"):
        identificar_formato(arquivos_cadastro[0].with_suffix(".csv").name, planilha, ("csv",))
    # A mesma planilha com um projeto de macros acrescentado
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(planilha)) as original, zipfile.ZipFile(buffer, "w") as nova:
        for item in original.infolist():
            nova.writestr(item, original.read(item))
        nova.writestr("xl/vbaProject.bin", planilha[:512])
    with pytest.raises(ErroImportacao, match="macros"):
        identificar_formato(arquivos_cadastro[0].name, buffer.getvalue(), ("xlsx",))


def test_mensagens_de_erro_nao_repetem_conteudo_real(tabelas_cadastro):
    emails = valores(tabelas_cadastro, "email", CAMPOS_CADASTRO)[:20]
    cpfs = [validar_cpf(c)[0] for c in valores(tabelas_cadastro, "cpf", CAMPOS_CADASTRO)[:20]]
    vazou = 0
    for email, cpf in zip(emails, cpfs):
        with pytest.raises(ErroImportacao) as erro:
            identificar_formato("dados.csv", f"{email};{cpf}\x00".encode(), ("csv",))
        vazou += str(email) in str(erro.value) or str(cpf) in str(erro.value)
    confere(vazou == 0, f"{vazou} mensagem(ns) de erro com conteúdo das células")
